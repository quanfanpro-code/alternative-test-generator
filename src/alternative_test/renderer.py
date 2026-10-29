"""沿用内置模板，按实际抽取行扩展底稿，并独立核对保存后的内容。"""

from __future__ import annotations

import copy
import json
import posixpath
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

from .errors import OutputConflictError
from .matching import _make_line_id
from .models import CounterpartyBalance, JournalLine, MatchDecision, ProcedureResult
from .controls import ControlResult
from .selection import SelectionResult
from .population import PopulationResult


# 参考 excel-master 的主题色；只保存配色，不依赖其他技能的安装位置。
THEMES = {
    "default": ("4472C4", "D9E2F3"), "deep-navy": ("1F4E79", "D6E4F0"),
    "jade": ("375623", "E2EFDA"), "slate": ("404040", "D9D9D9"),
    "burgundy": ("843C0C", "FCE4D6"), "coral": ("D84B4B", "FDE8E8"),
    "sakura": ("D94F70", "FDE8EE"), "warm-sun": ("E8843A", "FEF0E6"),
    "lavender": ("8B7EC8", "EDEAF8"), "matcha": ("7BA23F", "F0F5E6"),
    "peach": ("E8897C", "FEF0ED"), "misty": ("6B7FB5", "EDF0F8"),
}
AMOUNT_FMT = '#,##0.00;[Red]-#,##0.00'
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


@dataclass
class RenderContext:
    project_name: str = ""
    period: str = ""
    entity_ids: list[str] = field(default_factory=list)
    side: str = "liability"
    balances: list[CounterpartyBalance] = field(default_factory=list)
    journal_lines: list[JournalLine] = field(default_factory=list)
    match_decisions: list[MatchDecision] = field(default_factory=list)
    controls: list[ControlResult] = field(default_factory=list)
    selection: SelectionResult | None = None
    procedure_results: dict[str, ProcedureResult] = field(default_factory=dict)
    excluded_lines: list[tuple[JournalLine, str]] = field(default_factory=list)
    population: PopulationResult | None = None
    run_id: str = ""
    input_hashes: dict[str, str] = field(default_factory=dict)
    check_overrides: dict[str, list[dict]] = field(default_factory=dict)
    entity_names: dict[str, str] = field(default_factory=dict)
    period_start: date | None = None
    period_end: date | None = None
    post_period_end: date | None = None
    prepared_by: str = ""
    reviewed_by: str = ""
    prepared_date: date | None = None
    reviewed_date: date | None = None
    index_number: str = ""
    index_numbers: dict[str, str] = field(default_factory=dict)
    currency: str = "人民币"
    theme: str = "deep-navy"
    selection_mode: str = "population"  # 无规则时输出候选总体，不能称为已测试样本。


def _text(cell, value):
    """原始名称、摘要及凭证号始终作为文字，防止前导零丢失或误成公式。"""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _value(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _safe_title(raw, used):
    name = re.sub(r"[\\/*?:\[\]]", "_", raw).strip().strip("'") or "往来单位"
    stem = name[:31]
    title, n = stem, 1
    while title.casefold() in used:
        suffix = f"_{n}"
        title = stem[:31-len(suffix)] + suffix
        n += 1
    used.add(title.casefold())
    return title


def _period_end(ctx):
    if ctx.period_end:
        return ctx.period_end
    dates = re.findall(r"\d{4}-\d{2}-\d{2}", ctx.period)
    return date.fromisoformat(dates[-1]) if dates else None


def _locate_sections(ws):
    positions = []
    for row in ws:
        for cell in row:
            val = cell.value
            if isinstance(val, str) and re.match(r"^[一二三]、(?:借方发生额|贷方发生额|期后付款|期后收款)$", val):
                positions.append((cell.row, val))
    if len(positions) != 3:
        raise ValueError(f"工作表 {ws.title} 的三个凭证板块不完整，不能猜测填表位置")
    return positions


def _insert_rows(ws, at, count):
    """扩展空行同时移动行高；内置模板没有合并格及旧公式，公式在扩展后新建。"""
    if count <= 0:
        return
    dimensions = {r: copy.copy(d) for r, d in ws.row_dimensions.items() if r >= at}
    ws.insert_rows(at, count)
    for r in dimensions:
        del ws.row_dimensions[r]
    for r, dim in dimensions.items():
        dim.index = r + count
        ws.row_dimensions[r+count] = dim


def _matched(balance, ctx, post=False):
    ids = {d.journal_line_id for d in ctx.match_decisions
           if d.counterparty_id == balance.counterparty_id and d.status in ("exact_unique", "confirmed")}
    end = _period_end(ctx)
    source = [l for l, _ in ctx.excluded_lines] if post else ctx.journal_lines
    result = []
    seen = set()
    for line in source:
        key = _make_line_id(line)
        if key in seen or key not in ids or line.entity_id != balance.entity_id:
            continue
        if post:
            if end is None or line.posting_date <= end:
                continue
            if ctx.post_period_end and line.posting_date > ctx.post_period_end:
                continue
            if not post_payment_evidence(line, ctx.side)[0]:
                continue
        elif ((ctx.period_start and line.posting_date < ctx.period_start)
              or (end and line.posting_date > end)):
            continue
        seen.add(key)
        result.append(line)
    return result


def post_payment_evidence(line, side):
    """区分资金收付线索和非资金结转；候选不能等同于期末余额已清偿。"""
    direction = "debit" if side == "liability" else "credit"
    if getattr(line, direction) == 0:
        return False, "非期后收付款方向"
    account = line.opposing_account_name or ""
    if account:
        parts = re.split(r"[、；;]", account)
        cash = [bool(re.search(r"银行存款|库存现金|其他货币资金|(?:^|[- /])现金(?:$|[- /])", p)) for p in parts]
        if all(cash):
            return True, "资金科目收付款候选，仍需核对与测试时点余额的对应关系"
        if any(cash):
            return True, "混合对方科目，资金金额及余额对应关系待核查"
        return False, "对方科目为非资金科目，排除抵销或重分类等非收付款发生额"
    if re.search(r"抵销|抵消|重分类|重分|结转|转入|转出|冲账", line.summary):
        return False, "摘要显示抵销或结转，不能作为收付款"
    if re.search(r"付款|支付|收款|回款|收讫|付讫|汇款|银行", line.summary):
        return True, "摘要显示收付款线索；资金凭据及余额对应关系待核查"
    return False, "缺少收付款线索，保留在期后检索记录中，不冒充已收付款"


def _fill_info(ws, ctx, entity_id=None):
    if entity_id:
        name = ctx.entity_names.get(entity_id, entity_id)
    else:
        name = "、".join(ctx.entity_names.get(e, e) for e in ctx.entity_ids)
    for coord, val in {
        "C4": name, "F4": ctx.prepared_by, "I4": ctx.prepared_date,
        "L4": ctx.index_number, "C5": _period_end(ctx) or ctx.period,
        "F5": ctx.reviewed_by, "I5": ctx.reviewed_date, "L5": ctx.project_name,
    }.items():
        _text(ws[coord], val or None)
        if isinstance(val, date):
            ws[coord].number_format = "yyyy-mm-dd"
    return {c: _value(ws[c].value) for c in ("C4", "F4", "I4", "L4", "C5", "F5", "I5", "L5")}


def _style(ws, ctx, widths):
    for row in ws:
        for c in row:
            c.font = Font(name="微软雅黑", size=10, color="202020")
            c.alignment = RIGHT if isinstance(c.value, (int, float)) or c.data_type == "f" else LEFT
            c.border = Border()
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in range(1, ws.max_row+1):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[4].height = 32
    ws.row_dimensions[5].height = 32
    ws["B2"].font = Font(name="微软雅黑", size=15, bold=True, color=THEMES[ctx.theme][0])
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if ws.max_row <= 36 else 0
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.print_area = f"B2:M{ws.max_row}"
    ws["C4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    units = sum(2 if ord(ch) > 255 else 1 for ch in str(ws["C4"].value or ""))
    width = max(1, int(ws.column_dimensions["C"].width * 0.9))
    ws.row_dimensions[4].height = max(32, ((units+width-1)//width)*14+4)


def _band(ws, row, ctx, header=False, first=2, last=13):
    dark, pale = THEMES[ctx.theme]
    color = dark if header else pale
    for col in range(first, last+1):
        c = ws.cell(row, col)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF" if header else "202020")
        c.border = Border(bottom=Side(style="thin", color=dark))
    ws.row_dimensions[row].height = 24 if header else 20


def _fill_detail(ws, balance, ctx):
    info = _fill_info(ws, ctx, balance.entity_id)
    # I列在明细中是窄分隔列；日期放入J列，避免Excel显示井号。
    for source, destination in (("I4", "J4"), ("I5", "J5")):
        _text(ws[destination], ws[source].value)
        ws[destination].number_format = "yyyy-mm-dd"
        ws[source].value = None
        info[source] = None
        info[destination] = _value(ws[destination].value)
    _text(ws["B2"], f"{balance.account_name}替代测试候选底稿")
    _text(ws["B7"], "被测试往来单位：")
    _text(ws["C7"], balance.raw_name)
    _text(ws["L4"], ctx.index_numbers.get(balance.counterparty_id) or None)
    info["L4"] = ws["L4"].value
    for c, val in {"B8": "期初余额：", "C8": balance.opening_balance, "E8": "期末余额：",
                   "F8": balance.ending_balance, "H8": "币种：", "J8": ctx.currency}.items():
        _text(ws[c], _value(val))
    info.update({c: _value(ws[c].value) for c in ("B2", "C7", "C8", "F8", "J8")})
    all_lines = _matched(balance, ctx)
    chosen = {_make_line_id(l) for l, _ in ctx.selection.selected} if ctx.selection else set()
    selected = [l for l in all_lines if _make_line_id(l) in chosen] if ctx.selection_mode == "selected" else all_lines
    directions = ["credit", "debit", "debit"] if ctx.side == "liability" else ["debit", "credit", "credit"]
    post = _matched(balance, ctx, post=True)
    if ctx.selection_mode == "selected":
        post = [l for l in post if _make_line_id(l) in chosen]
    groups = [[l for l in selected if getattr(l, d) != 0] for d in directions[:2]]
    groups.append([l for l in post if getattr(l, directions[2]) != 0])
    original = _locate_sections(ws)
    # 从底部扩展，不让前一个板块的插入改变尚未处理的定位。
    for i in (2, 1, 0):
        start = original[i][0] + 3
        count = max(1, len(groups[i]))
        _insert_rows(ws, start+1, count-1 + (1 if i == 2 else 3))
    positions = _locate_sections(ws)
    sections, traces = [], []
    for i, ((heading, title), lines, direction) in enumerate(zip(positions, groups, directions)):
        start, count = heading+3, max(1, len(lines))
        end, subtotal = start+count-1, start+count
        for r in range(start, end+1):
            for col in range(2, 14):
                ws.cell(r, col).value = None
        for n, line in enumerate(lines, 1):
            r = start+n-1
            values = [n, line.posting_date, line.voucher_type, line.voucher_number,
                      line.summary, line.opposing_account_name, float(getattr(line, direction))]
            for col, val in enumerate(values, 2):
                _text(ws.cell(r, col), val)
            traces.append([ws.title, r, title, _make_line_id(line),
                           json.dumps([_value(v) for v in values], ensure_ascii=False),
                           json.dumps({"ledger": line.source.model_dump(mode="json"),
                               "opposing_accounts": [s.model_dump(mode="json") for s in line.opposing_account_sources]}, ensure_ascii=False)])
        ws.cell(subtotal, 2, "测试金额小计" if ctx.selection_mode == "selected" else "候选金额小计")
        ws.cell(subtotal, 8, f"=SUM(H{start}:H{end})")
        section = {"title": title, "heading": heading, "start": start, "end": end,
                   "subtotal": subtotal, "count": len(lines), "direction": direction,
                   "expected_total": str(sum((getattr(l, direction) for l in lines), Decimal(0)))}
        if i < 2:
            denominator, ratio = subtotal+1, subtotal+2
            annual = sum((getattr(l, direction) for l in all_lines), Decimal(0))
            ws.cell(denominator, 2, f"期间内{title.split('、', 1)[1]}合计")
            ws.cell(denominator, 8, float(annual))
            ws.cell(ratio, 2, "测试金额占期间发生额比例" if ctx.selection_mode == "selected" else "候选金额占期间发生额比例")
            formula = f'=IF(H{denominator}=0,"",H{subtotal}/H{denominator})'
            ws.cell(ratio, 8, formula)
            if annual == 0:
                ws.cell(ratio, 10, "期间发生额为0，比例不适用")
            section.update(denominator=denominator, annual=str(annual), ratio=ratio, ratio_formula=formula)
        sections.append(section)
    _style(ws, ctx, {"A": 2, "B": 23, "C": 16, "D": 12, "E": 14, "F": 38, "G": 22,
                     "H": 18, "I": 2, "J": 13, "K": 13, "L": 17, "M": 13, "N": 2})
    ws.freeze_panes = "F12"
    ws.print_title_rows = "2:5"
    ws.row_dimensions[7].height = 28
    ws["C7"].alignment = Alignment(horizontal="left", vertical="center")
    ws["C8"].number_format = ws["F8"].number_format = AMOUNT_FMT
    for sec in sections:
        _band(ws, sec["heading"], ctx, header=True)
        _band(ws, sec["heading"]+1, ctx)
        _band(ws, sec["heading"]+2, ctx, header=True)
        ws.row_dimensions[sec["heading"]+2].height = 30
        for r in range(sec["start"], sec["end"]+1):
            for col in range(2, 14):
                cell = ws.cell(r, col)
                cell.fill = PatternFill("solid", fgColor="F3F6F9" if (r-sec["start"]) % 2 else "FFFFFF")
                cell.border = Border(bottom=Side(style="dotted", color="C8D2DC"))
                cell.alignment = CENTER if col in (2, 3, 4, 5, 10, 12) else (RIGHT if col in (8, 11, 13) else LEFT)
            ws.cell(r, 3).number_format = "yyyy-mm-dd"
            ws.cell(r, 5).number_format = "@"
            ws.cell(r, 8).number_format = AMOUNT_FMT
            text = str(ws.cell(r, 6).value or "")
            units = sum(2 if ord(ch) > 255 else 1 for ch in text)
            ws.row_dimensions[r].height = max(20, ((units+31)//32)*16)
        _band(ws, sec["subtotal"], ctx)
        ws.cell(sec["subtotal"], 2).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(sec["subtotal"], 8).number_format = AMOUNT_FMT
        if "ratio" in sec:
            ws.cell(sec["denominator"], 8).number_format = AMOUNT_FMT
            ws.cell(sec["ratio"], 8).number_format = "0.00%"
            for row in (sec["denominator"], sec["ratio"]):
                ws.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[sec["ratio"]].height = 32
    ws.cell(ws.max_row, 2, "候选凭证不代表完成替代测试；期后收付款与测试时点余额的对应关系、附件和结论待核查。")
    ws.row_dimensions[ws.max_row].height = 32
    ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(ws.max_row, 2).font = Font(name="微软雅黑", size=9, color="596A78")
    return {"sheet": ws.title, "counterparty_id": balance.counterparty_id, "info": info, "sections": sections}, traces


def _fill_summary(ws, ctx):
    info = _fill_info(ws, ctx)
    count = len(ctx.balances)
    _insert_rows(ws, 8, count)
    expected = []
    for i, balance in enumerate(ctx.balances):
        r = 8+i
        ctrl = next((c for c in ctx.controls if c.counterparty_id == balance.counterparty_id), None)
        values = [ctx.index_numbers.get(balance.counterparty_id) or None, balance.raw_name, ctx.currency,
                  float(balance.ending_balance), None, "待核查" + (f"；余额勾稽：{ctrl.status.value}" if ctrl else "")]
        for col, val in enumerate(values, 2):
            _text(ws.cell(r, col), val)
        expected.append([r, values])
    total = 8+count
    ws.cell(total, 2, f"共 {count} 行")
    ws.cell(total, 5, f"=SUM(E8:E{total-1})" if count else "=0")
    ws.cell(total, 6, f'=IF(COUNT(F8:F{total-1})=0,"",SUM(F8:F{total-1}))' if count else '=""')
    _style(ws, ctx, {"A": 2, "B": 23, "C": 42, "D": 12, "E": 20, "F": 26, "G": 40,
                     "H": 10, "I": 14, "J": 2, "K": 12, "L": 18, "M": 3})
    _band(ws, 7, ctx, header=True, last=7)
    _band(ws, total, ctx, last=7)
    ws.row_dimensions[7].height = 38
    for r in range(8, total+1):
        ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = AMOUNT_FMT
        ws.row_dimensions[r].height = 30 if r != total else 22
    ws.freeze_panes = "D8"
    ws.cell(total+3, 2, "金额勾稽不代表已完成替代测试；可确认金额及审计结论待实际核查后填写。")
    ws.row_dimensions[total+3].height = 32
    ws.cell(total+3, 2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(total+3, 2).font = Font(name="微软雅黑", size=9, color="596A78")
    return {"sheet": ws.title, "info": info, "rows": expected, "total": total,
            "total_formula": ws.cell(total, 5).value, "confirmed_formula": ws.cell(total, 6).value}


def render_workbook(context, template, output):
    output, template = Path(output), Path(template)
    if output.exists() or output.resolve() == template.resolve():
        raise OutputConflictError(f"输出文件已存在或等于模板: {output}")
    if context.theme not in THEMES:
        raise ValueError(f"未知主题：{context.theme}")
    if context.procedure_results or context.check_overrides:
        raise ValueError("本生成器只写候选底稿；实际检查结果需使用经确认的独立核查流程")
    wb = openpyxl.load_workbook(template)
    try:
        required = ("替代结果汇总表", "替代测试表模板")
        if any(n not in wb.sheetnames for n in required):
            raise ValueError("模板缺少替代结果汇总表或替代测试表模板；本版本沿用内置双表模板")
        if any(s.merged_cells.ranges for s in wb):
            raise ValueError("内置模板不含合并格，不能直接套用其他结构的模板")
        if any(c.data_type == "f" for s in wb for row in s for c in row):
            raise ValueError("输入模板含旧公式，需先确认结构；不能猜测移动公式")
        _locate_sections(wb[required[1]])
        summary = _fill_summary(wb[required[0]], context)
        units, traces = [], []
        used = {n.casefold() for n in wb.sheetnames} | {"_运行信息", "_来源追踪"}
        for balance in context.balances:
            ws = wb.copy_worksheet(wb[required[1]])
            ws.title = _safe_title(balance.raw_name, used)
            unit, rows = _fill_detail(ws, balance, context)
            units.append(unit)
            traces.extend(rows)
        wb[required[1]].sheet_state = "hidden"
        meta = wb.create_sheet("_运行信息")
        meta.append(["schema", "alternative-test-render-v2"])
        meta.append(["summary", json.dumps(summary, ensure_ascii=False)])
        for unit in units:
            meta.append(["unit", json.dumps(unit, ensure_ascii=False)])
        meta.append(["run", json.dumps({"run_id": context.run_id, "theme": context.theme,
                    "selection_mode": context.selection_mode, "input_hashes": context.input_hashes}, ensure_ascii=False)])
        source = wb.create_sheet("_来源追踪")
        source.append(["工作表", "输出行", "板块", "凭证明细ID", "原始填写内容", "来源"])
        for row in traces:
            source.append(row)
        meta.sheet_state = source.sheet_state = "hidden"
        wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        buffer = BytesIO()
        wb.save(buffer)
    finally:
        wb.close()
    report = validate_workbook(BytesIO(buffer.getvalue()))
    if not report["valid"]:
        raise ValueError("生成前验证失败：" + "；".join(report["errors"]))
    output.parent.mkdir(parents=True, exist_ok=True)
    # 独占创建，既不覆盖已有文件，也不创建需要永久删除的临时目录。
    with output.open("xb") as stream:
        stream.write(buffer.getvalue())
    report = validate_workbook(output)
    if not report["valid"]:
        raise ValueError("保存后验证失败，保留文件供排查：" + "；".join(report["errors"]))
    return output


def validate_workbook(path):
    """重开文件核对来源、行数、公式、金额、空检查栏及 XML 声明，不冒充 Excel 重算。"""
    errors, checks = [], 0
    raw = path.getvalue() if isinstance(path, BytesIO) else Path(path).read_bytes()
    try:
        with ZipFile(BytesIO(raw)) as package:
            if package.testzip():
                errors.append("工作簿压缩包校验失败")
            for name in package.namelist():
                if name.endswith(".rels"):
                    root = ET.fromstring(package.read(name))
                    base = name.rsplit("/_rels/", 1)[0] if "/_rels/" in name else ""
                    for relation in root:
                        if relation.get("TargetMode") == "External":
                            continue
                        target = relation.get("Target", "")
                        resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                        if resolved not in package.namelist():
                            errors.append(f"XML 引用关系缺失：{name} -> {target}")
                if not name.endswith(".xml"):
                    continue
                scopes = [{}]
                pending = {}
                for event, obj in ET.iterparse(BytesIO(package.read(name)), events=("start-ns", "start", "end")):
                    if event == "start-ns":
                        pending[obj[0]] = obj[1]
                    elif event == "start":
                        scope = {**scopes[-1], **pending}
                        pending.clear()
                        scopes.append(scope)
                        ignorable = obj.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "")
                        for prefix in ignorable.split():
                            if prefix not in scope:
                                errors.append(f"XML 命名空间未声明：{name}：{prefix}")
                    else:
                        scopes.pop()
                checks += 1
        wb = openpyxl.load_workbook(BytesIO(raw))
        try:
            if "_运行信息" not in wb.sheetnames or "_来源追踪" not in wb.sheetnames:
                raise ValueError("缺少生成元数据，不能把普通空模板视为已验证成品")
            meta = list(wb["_运行信息"].values)
            if meta[0][1] != "alternative-test-render-v2":
                raise ValueError("运行信息版本不支持")
            summary = next(json.loads(v) for k, v in meta if k == "summary")
            units = [json.loads(v) for k, v in meta if k == "unit"]
            expected = {}
            for sheet, row, section, line_id, values, source in list(wb["_来源追踪"].values)[1:]:
                key = (sheet, row)
                if key in expected:
                    errors.append(f"来源追踪行重复：{sheet}!{row}")
                expected[key] = json.loads(values)
                json.loads(source)
            def same(sheet, coord, wanted):
                nonlocal checks
                checks += 1
                actual = _value(wb[sheet][coord].value)
                if isinstance(wanted, (int, float)) and isinstance(actual, (int, float)):
                    equal = abs(Decimal(str(actual))-Decimal(str(wanted))) <= Decimal("0.000001")
                else:
                    equal = actual == wanted or (actual is None and wanted == "")
                if not equal:
                    errors.append(f"数据/公式不一致：{sheet}!{coord}，应为{wanted!r}，实际{actual!r}")
            for unit in [summary, *units]:
                name = unit["sheet"]
                if name not in wb.sheetnames:
                    errors.append(f"缺少单位工作表：{name}")
                    continue
                for coord, val in unit["info"].items():
                    same(name, coord, val)
                if wb[name].sheet_view.showGridLines or wb[name].freeze_panes is None:
                    errors.append(f"成品格式缺失：{name}")
            sname = summary["sheet"]
            for row, values in summary["rows"]:
                for col, val in enumerate(values, 2):
                    same(sname, f"{openpyxl.utils.get_column_letter(col)}{row}", val)
            same(sname, f'E{summary["total"]}', summary["total_formula"])
            same(sname, f'F{summary["total"]}', summary["confirmed_formula"])
            expected_keys = set()
            for unit in units:
                name = unit["sheet"]
                if name not in wb.sheetnames:
                    continue
                for sec in unit["sections"]:
                    same(name, f'B{sec["heading"]}', sec["title"])
                    start, end, total = sec["start"], sec["end"], sec["subtotal"]
                    same(name, f"H{total}", f"=SUM(H{start}:H{end})")
                    actual_total = Decimal(0)
                    for r in range(start, end+1):
                        key = (name, r)
                        values = expected.get(key, [None]*7)
                        if key in expected:
                            expected_keys.add(key)
                        for col, val in enumerate(values, 2):
                            same(name, f"{openpyxl.utils.get_column_letter(col)}{r}", val)
                        amount = wb[name].cell(r, 8).value
                        if isinstance(amount, (int, float)):
                            actual_total += Decimal(str(amount))
                        for col in range(10, 14):
                            same(name, f"{openpyxl.utils.get_column_letter(col)}{r}", None)
                    if len([k for k in expected if k[0] == name and start <= k[1] <= end]) != sec["count"]:
                        errors.append(f"凭证行数不一致：{name}，{sec['title']}")
                    if abs(actual_total-Decimal(sec["expected_total"])) > Decimal("0.000001"):
                        errors.append(f"金额小计不一致：{name}，{sec['title']}")
                    if "ratio" in sec:
                        same(name, f'H{sec["denominator"]}', float(sec["annual"]))
                        same(name, f'H{sec["ratio"]}', sec["ratio_formula"])
                        if wb[name].cell(sec["ratio"], 8).number_format != "0.00%":
                            errors.append(f"比例格式错误：{name}")
            if expected_keys != set(expected):
                errors.append("来源追踪存在未落入板块的凭证")
            for sheet in wb:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == "f" and ("#REF!" in str(cell.value) or "[" in str(cell.value)):
                            errors.append(f"错误或外部公式引用：{sheet.title}!{cell.coordinate}")
                        if cell.data_type == "e":
                            errors.append(f"公式错误缓存：{sheet.title}!{cell.coordinate}")
        finally:
            wb.close()
    except Exception as exc:
        errors.append(f"验证无法完成：{exc}")
    return {"valid": not errors, "checks": checks, "errors": errors,
            "formula_validation": "公式范围与来源金额已独立核对；Excel 打开时重算"}
