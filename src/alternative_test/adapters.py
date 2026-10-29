"""表格适配器 — 列映射与行读取。"""

from __future__ import annotations

import csv
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import xlrd

from .fileio import detect_format, probe_file  # noqa: F401

# ──────────────────────────────────────────────
# 列映射
# ──────────────────────────────────────────────

@dataclass
class ColumnMapping:
    """列名 → 实际表头列名映射。None 表示该字段不可用。"""

    date_col: str | None = None
    voucher_type_col: str | None = None
    voucher_no_col: str | None = None
    summary_col: str | None = None
    account_col: str | None = None
    opposing_account_col: str | None = None
    ledger_col: str | None = None
    debit_col: str | None = None
    credit_col: str | None = None
    counterparty_aux_col: str | None = None

    # 额外的余额表字段
    entity_col: str | None = None
    counterparty_col: str | None = None
    opening_balance_col: str | None = None
    ending_balance_col: str | None = None
    account_code_col: str | None = None


def _build_col_index(headers: list[str]) -> dict[str, int]:
    """构建表头名→列索引映射。"""
    return {h: i for i, h in enumerate(headers)}


def _map_row(
    row: list[Any],
    headers: list[str],
    col_index: dict[str, int],
    mapping: ColumnMapping,
) -> dict[str, Any]:
    """按映射关系取出行中各字段值。"""
    result: dict[str, Any] = {}
    field_map = {
        "日期": mapping.date_col,
        "凭证类型": mapping.voucher_type_col,
        "凭证号": mapping.voucher_no_col,
        "摘要": mapping.summary_col,
        "科目": mapping.account_col,
        "对方科目": mapping.opposing_account_col,
        "账套": mapping.ledger_col,
        "借方": mapping.debit_col,
        "贷方": mapping.credit_col,
        "辅助核算": mapping.counterparty_aux_col,
        "公司": mapping.entity_col,
        "供应商": mapping.counterparty_col,
        "期初余额": mapping.opening_balance_col,
        "期末余额": mapping.ending_balance_col,
        "科目代码": mapping.account_code_col,
    }
    for key, col_name in field_map.items():
        if col_name and col_name in col_index:
            idx = col_index[col_name]
            result[key] = row[idx] if idx < len(row) else None
        else:
            result[key] = None
    return result


# ──────────────────────────────────────────────
# 读取接口
# ──────────────────────────────────────────────

def read_rows(
    path: Path,
    mapping: ColumnMapping,
    sheet: str | None = None,
    header_row: int = 1,
) -> Iterator[dict[str, Any]]:
    """根据列映射读取文件中的行（跳过表头）。"""
    fmt = detect_format(path)

    if fmt == "xlsx":
        yield from _read_xlsx(path, mapping, sheet, header_row)
    elif fmt == "xls":
        yield from _read_xls(path, mapping, sheet, header_row)
    else:
        yield from _read_csv(path, mapping, header_row)


def _read_xlsx(
    path: Path,
    mapping: ColumnMapping,
    sheet: str | None = None,
    header_row: int = 1,
) -> Iterator[dict[str, Any]]:
    """读取 xlsx 文件。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if sheet:
        ws = wb[sheet]
    else:
        ws = wb[wb.sheetnames[0]]

    # 读表头
    first_row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), None)
    headers = [str(c) if c is not None else "" for c in (first_row or [])]
    col_index = _build_col_index(headers)

    # 读数据行
    for row_number, row in enumerate(ws.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
        row_vals = list(row)
        result = _map_row(row_vals, headers, col_index, mapping)
        result["_source_row"] = row_number
        yield result

    wb.close()


def _read_xls(
    path: Path,
    mapping: ColumnMapping,
    sheet: str | None = None,
    header_row: int = 1,
) -> Iterator[dict[str, Any]]:
    """读取 xls 文件。"""
    wb = xlrd.open_workbook(str(path), on_demand=True)
    if sheet:
        ws = wb.sheet_by_name(sheet)
    else:
        ws = wb.sheet_by_index(0)

    # 读表头
    headers = [str(ws.cell_value(header_row-1, c)) for c in range(ws.ncols)] if ws.nrows >= header_row else []
    col_index = _build_col_index(headers)

    # 读数据行
    for r in range(header_row, ws.nrows):
        row_vals = [ws.cell_value(r, c) for c in range(ws.ncols)]
        result = _map_row(row_vals, headers, col_index, mapping)
        if mapping.date_col and mapping.date_col in col_index:
            ci = col_index[mapping.date_col]
            if ws.cell_type(r, ci) == xlrd.XL_CELL_DATE:
                result["日期"] = xlrd.xldate_as_datetime(ws.cell_value(r, ci), wb.datemode)
        result["_source_row"] = r+1
        yield result

    wb.release_resources()


def _read_csv(
    path: Path,
    mapping: ColumnMapping,
    header_row: int = 1,
) -> Iterator[dict[str, Any]]:
    """读取 CSV 文件（自动检测编码和分隔符）。"""
    # 检测编码
    for encoding in ["utf-8-sig", "utf-8", "gb18030", "gbk", "gb2312"]:
        try:
            with open(path, encoding=encoding) as f:
                sample = f.read()
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        encoding = "utf-8"
        with open(path, encoding=encoding, errors="replace") as f:
            sample = f.read()

    # 检测分隔符
    tabs = sample.count("\t")
    commas = sample.count(",")
    delimiter = "\t" if tabs > commas else ","

    with open(path, encoding=encoding) as f:
        reader = csv.reader(f, delimiter=delimiter)
        for _ in range(header_row-1):
            next(reader, None)
        headers = next(reader, None)
        if not headers:
            return
        headers = [h.strip().strip('"') for h in headers]
        col_index = _build_col_index(headers)

        for row_number, row in enumerate(reader, header_row+1):
            result = _map_row(row, headers, col_index, mapping)
            result["_source_row"] = row_number
            yield result


# ──────────────────────────────────────────────
# 适配器查找（为后续 CLI 做准备）
# ──────────────────────────────────────────────

def get_adapter(path: Path) -> str:
    """返回适配器名称。"""
    return detect_format(path)
