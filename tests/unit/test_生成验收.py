"""验证实际填表、扩展行、公式、数据隔离和正式生成入口。"""

import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from alternative_test.matching import propose_matches
from alternative_test.models import CounterpartyBalance, JournalLine, NormalBalance, SourceRef
from alternative_test.renderer import RenderContext, render_workbook
from alternative_test.selection import SelectionResult


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "assets/templates/替代测试-负债方-v1.xlsx"


def context(count=3):
    source = SourceRef(file_id="演示账簿.xlsx", sha256="a" * 64, sheet="账簿", row=2)
    balance = CounterpartyBalance(
        entity_id="E1", counterparty_id="CP1", raw_name="演示供应商有限公司",
        account_name="应付账款", normal_balance=NormalBalance.CREDIT,
        opening_balance=Decimal("100"), ending_balance=Decimal("100"), source=source,
    )
    lines = [JournalLine(
        entity_id="E1", ledger_id="L1", posting_date=date(2025, 6, 1),
        voucher_type="记", voucher_number=f"000{i+1}", line_number=i+1,
        account_name="应付账款", counterparty_raw=balance.raw_name,
        summary="演示采购", debit=Decimal("0"),
        credit=Decimal("-10") if i == 1 else Decimal("100"), source=source,
    ) for i in range(count)]
    return RenderContext(
        project_name="演示项目", period="2025-01-01 ~ 2025-12-31",
        entity_ids=["E1"], balances=[balance], journal_lines=lines,
        match_decisions=propose_matches(lines, [balance]).decisions,
    )


@pytest.mark.parametrize("count", [0, 1, 3, 50])
def test_inserted_rows_formula_and_data(tmp_path, count):
    from alternative_test.renderer import validate_workbook
    ctx = context(count)
    ctx.entity_names = {"E1": "演示被审计单位有限公司"}
    ctx.period_end = date(2025, 12, 31)
    result = render_workbook(ctx, TEMPLATE, tmp_path / "底稿.xlsx")
    report = validate_workbook(result)
    assert report["valid"], report
    w = openpyxl.load_workbook(result)
    s = w[ctx.balances[0].raw_name]
    assert s["C4"].value == "演示被审计单位有限公司"
    assert s["C7"].value == ctx.balances[0].raw_name
    data_end = 12 + max(count, 1) - 1
    assert s.cell(data_end+1, 8).value == f"=SUM(H12:H{data_end})"
    assert s.cell(data_end+2, 8).value == sum(float(l.credit) for l in ctx.journal_lines)
    assert s.cell(data_end+3, 8).data_type == "f"
    if count:
        assert s["E12"].value == "0001"
    if count > 1:
        assert s["H13"].value == -10
    assert any(s.cell(r, 2).value == "二、借方发生额" for r in range(1, s.max_row+1))
    assert not s.sheet_view.showGridLines
    assert not s.merged_cells.ranges
    w.close()


def test_selection_only_and_annual_denominator(tmp_path):
    ctx = context(3)
    ctx.selection_mode = "selected"
    ctx.selection = SelectionResult(population=ctx.journal_lines, selected=[(ctx.journal_lines[0], "R1")])
    result = render_workbook(ctx, TEMPLATE, tmp_path / "抽取.xlsx")
    w = openpyxl.load_workbook(result)
    s = w[ctx.balances[0].raw_name]
    assert s["E12"].value == "0001"
    assert s["B13"].value == "测试金额小计"
    assert s["H14"].value == 190
    assert s["B16"].value == "二、借方发生额"
    assert w["替代结果汇总表"]["F8"].value is None
    w.close()


def test_validator_detects_amount_and_formula_damage(tmp_path):
    from alternative_test.renderer import validate_workbook
    result = render_workbook(context(), TEMPLATE, tmp_path / "完整.xlsx")
    w = openpyxl.load_workbook(result)
    s = w["演示供应商有限公司"]
    s["H12"] = 9999
    s["H15"] = "=SUM(H12:H12)"
    damaged = tmp_path / "损坏副本.xlsx"
    w.save(damaged)
    w.close()
    report = validate_workbook(damaged)
    assert not report["valid"]
    assert any("金额" in e or "数据" in e for e in report["errors"])
    assert any("公式" in e for e in report["errors"])


def test_full_cli_pipeline(tmp_path):
    from alternative_test.cli import cmd_inspect, cmd_prepare, cmd_generate, cmd_validate
    balance = tmp_path / "余额.csv"
    journal = tmp_path / "序时账.csv"
    balance.write_text("供应商,期初余额,期末余额\n演示供应商有限公司,100,270\n", encoding="utf-8-sig")
    journal.write_text("日期,凭证号,供应商,借方,贷方,摘要,科目,对方科目\n"
        "2025-06-01,0001,演示供应商有限公司,0,100,演示采购,应付账款,原材料\n"
        "2025-07-01,0002,演示供应商有限公司,0,100,演示采购,应付账款,原材料\n"
        "2025-08-01,0003,演示供应商有限公司,30,0,演示付款,应付账款,银行存款\n"
        "2026-01-01,0004,演示供应商有限公司,50,0,期后演示付款,应付账款,银行存款\n", encoding="utf-8-sig")
    answers = {
        "project_name": "演示项目", "period_start": "2025-01-01", "period_end": "2025-12-31",
        "entity_ids": ["E1"], "entity_names": {"E1": "演示被审计单位有限公司"},
        "normal_balance": "credit", "prepared_by": "演示编制人",
        "inputs": [
            {"role": "balance", "path": str(balance), "entity_id": "E1", "account_name": "应付账款",
             "columns": {"counterparty_col": "供应商", "opening_balance_col": "期初余额", "ending_balance_col": "期末余额"}},
            {"role": "journal", "path": str(journal), "entity_id": "E1", "ledger_id": "L1",
             "columns": {"date_col": "日期", "voucher_no_col": "凭证号", "counterparty_aux_col": "供应商",
                         "debit_col": "借方", "credit_col": "贷方", "summary_col": "摘要", "account_col": "科目",
                         "opposing_account_col": "对方科目"}},
        ],
        "rules": [{"rule_id": "R1", "field": "absolute_amount", "operator": "gte", "value": "50", "user_confirmed": True}],
    }
    answer_path = tmp_path / "确认.json"
    answer_path.write_text(json.dumps(answers, ensure_ascii=False), encoding="utf-8")
    inspect_path = tmp_path / "探测.json"
    for func, args in [
        (cmd_inspect, ([str(balance), str(journal)], str(inspect_path))),
        (cmd_prepare, (str(inspect_path), str(answer_path), str(tmp_path / "准备"))),
        (cmd_generate, (str(tmp_path / "准备"), str(tmp_path / "成品"))),
        (cmd_validate, (str(tmp_path / "成品"),)),
    ]:
        with pytest.raises(SystemExit) as exc:
            func(*args)
        assert exc.value.code == 0
    outputs = list((tmp_path / "成品").glob("*.xlsx"))
    assert len(outputs) == 1
    w = openpyxl.load_workbook(outputs[0])
    s = w["演示供应商有限公司"]
    assert s["G12"].value == "原材料"
    assert w["替代结果汇总表"]["F8"].value is None
    values = [c.value for row in s for c in row]
    assert "0003" not in values
    assert "0004" in values
    w.close()


@pytest.mark.parametrize("side", ["liability", "asset"])
def test_multiyear_post_cash_and_non_cash(tmp_path, side):
    ctx = context(1)
    ctx.side = side
    ctx.period_start, ctx.period_end = date(2025, 1, 1), date(2025, 12, 31)
    ctx.post_period_end = date(2027, 12, 31)
    base = ctx.journal_lines[0]
    post = []
    for n, (year, opposing, description) in enumerate([
        (2024, "银行存款", "以前年度付款"), (2026, "银行存款", "期后付款"),
        (2027, "银行存款", "期后付款"), (2026, "其他应付款", "抵销重分类"),
        (2028, "银行存款", "范围外付款"),
    ], 10):
        post.append(base.model_copy(update={"posting_date": date(year, 1, 1), "line_number": n,
            "voucher_number": f"期后{n}", "opposing_account_name": opposing, "summary": description,
            "debit": Decimal("50") if side == "liability" else Decimal("0"),
            "credit": Decimal("0") if side == "liability" else Decimal("50")}))
    if side == "asset":
        ctx.balances[0].normal_balance = NormalBalance.DEBIT
        base.debit, base.credit = base.credit, base.debit
    ctx.excluded_lines = [(l, "期间外") for l in post]
    ctx.match_decisions = propose_matches([base, *post], ctx.balances).decisions
    filename = "替代测试-负债方-v1.xlsx" if side == "liability" else "替代测试-资产方-v1.xlsx"
    result = render_workbook(ctx, ROOT / "assets/templates" / filename, tmp_path / "期后.xlsx")
    w = openpyxl.load_workbook(result)
    s = w[ctx.balances[0].raw_name]
    values = [c.value for row in s for c in row]
    assert "期后11" in values and "期后12" in values
    assert all(number not in values for number in ("期后10", "期后13", "期后14"))
    w.close()


def test_no_selected_vouchers_stays_blank(tmp_path):
    ctx = context(3)
    ctx.selection_mode = "selected"
    ctx.selection = SelectionResult(population=ctx.journal_lines)
    output = render_workbook(ctx, TEMPLATE, tmp_path / "未抽中.xlsx")
    w = openpyxl.load_workbook(output)
    s = w[ctx.balances[0].raw_name]
    assert s["E12"].value is None
    assert s["H14"].value == 190
    w.close()


def test_same_name_different_entities_do_not_mix(tmp_path):
    ctx = context(1)
    ctx.entity_ids = ["E1", "E2"]
    ctx.entity_names = {"E1": "演示甲单位", "E2": "演示乙单位"}
    second = ctx.balances[0].model_copy(update={"entity_id": "E2", "counterparty_id": "CP2"})
    second_line = ctx.journal_lines[0].model_copy(update={"entity_id": "E2", "credit": Decimal("777")})
    ctx.balances.append(second)
    ctx.journal_lines.append(second_line)
    ctx.match_decisions = propose_matches(ctx.journal_lines, ctx.balances).decisions
    output = render_workbook(ctx, TEMPLATE, tmp_path / "同名.xlsx")
    w = openpyxl.load_workbook(output)
    assert w["演示供应商有限公司"]["H12"].value == 100
    assert w["演示供应商有限公司_1"]["H12"].value == 777
    assert w["演示供应商有限公司_1"]["C4"].value == "演示乙单位"
    w.close()


def test_long_invalid_names_and_formula_text_are_preserved(tmp_path):
    ctx = context(1)
    raw = "很长的演示供应商名称" * 5 + "[甲]/公司"
    ctx.balances[0].raw_name = raw
    ctx.journal_lines[0].counterparty_raw = raw
    ctx.journal_lines[0].summary = "=只是原始摘要，不能作为公式"
    ctx.match_decisions = propose_matches(ctx.journal_lines, ctx.balances).decisions
    output = render_workbook(ctx, TEMPLATE, tmp_path / "长名称.xlsx")
    w = openpyxl.load_workbook(output)
    s = next(s for s in w if s.sheet_state == "visible" and s.title != "替代结果汇总表")
    assert len(s.title) <= 31
    assert s["C7"].value == raw
    assert s["F12"].value == ctx.journal_lines[0].summary
    assert s["F12"].data_type == "s"
    w.close()


def test_zero_denominator_has_explicit_guard(tmp_path):
    ctx = context(2)
    ctx.journal_lines[1].credit = Decimal("-100")
    output = render_workbook(ctx, TEMPLATE, tmp_path / "零分母.xlsx")
    w = openpyxl.load_workbook(output)
    s = w[ctx.balances[0].raw_name]
    assert s["H15"].value == 0
    assert s["H16"].value == '=IF(H15=0,"",H14/H15)'
    assert s["J16"].value == "期间发生额为0，比例不适用"
    w.close()


def test_signed_rules_do_not_take_absolute_value():
    from alternative_test.population import PopulationResult
    from alternative_test.selection import SelectionRule, select_candidates
    ctx = context(2)
    result = select_candidates(PopulationResult(included=ctx.journal_lines), [SelectionRule(
        rule_id="R1", field="signed_amount", operator="lte", value="-5", user_confirmed=True)])
    assert [l.credit for l, _ in result.selected] == [Decimal("-10")]


def test_validator_catches_undeclared_ignorable_prefix(tmp_path):
    from io import BytesIO
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile
    from alternative_test.renderer import validate_workbook
    output = render_workbook(context(1), TEMPLATE, tmp_path / "正确.xlsx")
    damaged = BytesIO()
    with ZipFile(output) as original, ZipFile(damaged, "w") as changed:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                root.set("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "undeclared")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            changed.writestr(name, data)
    report = validate_workbook(damaged)
    assert not report["valid"]
    assert any("命名空间未声明" in e for e in report["errors"])
