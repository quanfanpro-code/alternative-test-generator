"""渲染回归：两类科目、真实数据、空检查栏以及覆盖保护。"""
from pathlib import Path
import openpyxl
import pytest
from alternative_test.errors import OutputConflictError
from alternative_test.matching import propose_matches
from alternative_test.models import NormalBalance
from alternative_test.renderer import render_workbook, validate_workbook
from tests.unit.test_生成验收 import context, ROOT

@pytest.mark.parametrize("side", ["liability", "asset"])
def test_both_sides_contain_actual_vouchers(tmp_path, side):
    ctx = context(3)
    ctx.side = side
    if side == "asset":
        ctx.balances[0].normal_balance = NormalBalance.DEBIT
        ctx.balances[0].account_name = "应收账款"
        for line in ctx.journal_lines:
            line.account_name = "应收账款"
            line.debit, line.credit = line.credit, line.debit
        ctx.match_decisions = propose_matches(ctx.journal_lines, ctx.balances).decisions
    filename = "替代测试-负债方-v1.xlsx" if side == "liability" else "替代测试-资产方-v1.xlsx"
    result = render_workbook(ctx, ROOT / "assets/templates" / filename, tmp_path / "底稿.xlsx")
    assert validate_workbook(result)["valid"]
    w = openpyxl.load_workbook(result)
    s = w[ctx.balances[0].raw_name]
    assert s["E12"].value == "0001"
    assert s["H13"].value == -10
    assert "贷方发生额" in s["B9"].value if side == "liability" else "借方发生额" in s["B9"].value
    for r in range(12, 15):
        assert all(s.cell(r, c).value is None for c in range(10, 14))
    assert w["替代结果汇总表"]["F8"].value is None
    w.close()

def test_conflict_preserves_existing_file(tmp_path):
    output = tmp_path / "不能覆盖.xlsx"
    output.write_bytes(b"existing")
    with pytest.raises(OutputConflictError):
        render_workbook(context(), ROOT / "assets/templates/替代测试-负债方-v1.xlsx", output)
    assert output.read_bytes() == b"existing"

def test_no_automatic_checkmarks_or_conclusions(tmp_path):
    output = render_workbook(context(), ROOT / "assets/templates/替代测试-负债方-v1.xlsx", tmp_path / "空检查.xlsx")
    w = openpyxl.load_workbook(output)
    for s in w:
        for row in s:
            for c in row:
                assert all(word not in str(c.value) for word in ("√", "未发现异常", "附件完整"))
    w.close()
